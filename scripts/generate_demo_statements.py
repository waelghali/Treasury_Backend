import os
import sys
from decimal import Decimal
from datetime import datetime, timedelta
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure project root in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from app.database import SessionLocal, engine, Base
from app.models.models_reconciliation_v2 import InternalLedgerRecord, BankStatement, BankTransaction
from app.models import Customer, Bank

BANKS_DATA = [
    {"name": "HSBC Bank Egypt S.A.E.", "code": "HSBC", "bank_id": 1},
    {"name": "Commercial International Bank (CIB)", "code": "CIB", "bank_id": 2},
    {"name": "National Bank of Egypt (NBE)", "code": "NBE", "bank_id": 3},
    {"name": "QNB Alahli", "code": "QNB", "bank_id": 4},
    {"name": "Banque Misr", "code": "MISR", "bank_id": 5},
    {"name": "First Abu Dhabi Bank (FAB Misr)", "code": "FAB", "bank_id": 6},
    {"name": "Citibank Egypt", "code": "CITI", "bank_id": 7},
    {"name": "Arab African International Bank (AAIB)", "code": "AAIB", "bank_id": 8},
    {"name": "Bank of Alexandria (Intesa Sanpaolo)", "code": "ALEX", "bank_id": 9},
    {"name": "Abu Dhabi Commercial Bank (ADCB Egypt)", "code": "ADCB", "bank_id": 10},
    {"name": "Attijariwafa Bank Egypt", "code": "AWB", "bank_id": 11},
    {"name": "Emirates NBD Egypt", "code": "ENBD", "bank_id": 12},
    {"name": "Standard Chartered Bank", "code": "SCB", "bank_id": 13},
    {"name": "Credit Agricole Egypt", "code": "CAE", "bank_id": 14},
    {"name": "Faisal Islamic Bank of Egypt", "code": "FAISAL", "bank_id": 15},
    {"name": "Al Ahli Bank of Kuwait (ABK Egypt)", "code": "ABK", "bank_id": 16},
    {"name": "Midbank", "code": "MID", "bank_id": 17},
    {"name": "Suez Canal Bank", "code": "SCB_EG", "bank_id": 18},
    {"name": "Arab Bank Egypt", "code": "ARAB", "bank_id": 19},
    {"name": "Housing & Development Bank (HDB)", "code": "HDB", "bank_id": 20},
]

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'demo_statements'))
os.makedirs(OUTPUT_DIR, exist_ok=True)

def generate_statements_and_erp_records():
    db = SessionLocal()
    
    # Verify customer exists
    customer = db.query(Customer).first()
    customer_id = customer.id if customer else 1
    print(f"Using customer_id: {customer_id}")

    # Clear previous simulated demo ERP records to stay clean
    db.query(InternalLedgerRecord).filter(
        InternalLedgerRecord.company_id == customer_id,
        InternalLedgerRecord.reference_number.like('DEMO-%')
    ).delete(synchronize_session=False)
    db.commit()

    generated_files = []
    total_erp_records_created = 0

    base_date = datetime(2026, 1, 1)

    for i in range(1, 21):
        bank_info = BANKS_DATA[(i - 1) % len(BANKS_DATA)]
        bank_name = bank_info["name"]
        bank_code = bank_info["code"]
        
        # Unique monthly sequence
        month_offset = (i - 1)
        # Advance by months cleanly
        stmt_year = 2026 + (month_offset // 12)
        stmt_month = (month_offset % 12) + 1
        stmt_start = datetime(stmt_year, stmt_month, 1)
        
        # End of month
        if stmt_month in [1, 3, 5, 7, 8, 10, 12]:
            end_day = 31
        elif stmt_month in [4, 6, 9, 11]:
            end_day = 30
        else:
            end_day = 28
        stmt_end = datetime(stmt_year, stmt_month, end_day)

        account_no = f"001-987654-{i:03d}"
        opening_bal = Decimal(f"{1500000.00 + (i * 125000.00):.2f}")
        current_bal = opening_bal

        # Build Transactions for this statement
        # 1. Bank Charge: Monthly fee
        # 2. Bank Charge: Stamp duty
        # 3. Bank Charge: SWIFT transfer
        # 4. Interest Earned
        # 5. Tax on interest
        # 6-7. Exact Sweep Pair
        # 8-9. Sweep with Fee Variance Pair
        # 10-11. Error & Reversal Pair
        # 12-14. Customer Collections (AR)
        # 15-16. Supplier Disbursements (AP)
        # 17. Corporate Payroll
        # 18. LG Issuance Commission
        # 19. Narrative Month Token
        # 20. Cryptic Memo

        raw_txns = [
            {
                "day": 2, "desc": "MONTHLY CORPORATE ACCOUNT MAINTENANCE FEE",
                "debit": Decimal("250.00"), "credit": Decimal("0.00"), "ref": f"CHG-{i:02d}01"
            },
            {
                "day": 3, "desc": "STATUTORY FISCAL STAMP DUTY CHARGE TAX LAW",
                "debit": Decimal("65.00"), "credit": Decimal("0.00"), "ref": f"STAMP-{i:02d}02"
            },
            {
                "day": 4, "desc": "SWIFT MT103 OUTWARD TRANSFER COMMISSION CHARGES",
                "debit": Decimal("350.00"), "credit": Decimal("0.00"), "ref": f"SWF-{i:02d}03"
            },
            {
                "day": 5, "desc": "CREDIT INTEREST EARNED ON OVERNIGHT DEPOSIT",
                "debit": Decimal("0.00"), "credit": Decimal("18500.00"), "ref": f"INT-{i:02d}04"
            },
            {
                "day": 5, "desc": "TAX DEDUCTION 20% WITHHOLDING TAX ON DEPOSIT INTEREST",
                "debit": Decimal("3700.00"), "credit": Decimal("0.00"), "ref": f"WHT-{i:02d}05"
            },
            # Exact Sweep Pair
            {
                "day": 8, "desc": f"SWEEP TO EG {bank_code} 001-999-001 GLE{i:02d}101_SWEEP",
                "debit": Decimal("400000.00"), "credit": Decimal("0.00"), "ref": f"GLE{i:02d}101"
            },
            {
                "day": 8, "desc": f"SWEEP FROM EG {bank_code} 001-999-001 GLE{i:02d}101_SWEEP",
                "debit": Decimal("0.00"), "credit": Decimal("400000.00"), "ref": f"GLE{i:02d}101"
            },
            # Sweep with Fee Variance (Outflow -250,075 vs Inflow +250,000)
            {
                "day": 10, "desc": f"CONCENTRATION POOLING SWEEP OUTFLOW GLE{i:02d}202_POOL",
                "debit": Decimal("250075.00"), "credit": Decimal("0.00"), "ref": f"GLE{i:02d}202"
            },
            {
                "day": 10, "desc": f"CONCENTRATION POOLING SWEEP INFLOW GLE{i:02d}202_POOL",
                "debit": Decimal("0.00"), "credit": Decimal("250000.00"), "ref": f"GLE{i:02d}202"
            },
            # Error & Reversal Pair
            {
                "day": 12, "desc": f"DUPLICATE DEBIT BANK ERROR REF ERR{i:02d}99",
                "debit": Decimal("14250.00"), "credit": Decimal("0.00"), "ref": f"ERR{i:02d}99"
            },
            {
                "day": 12, "desc": f"CORRECTION REVERSAL OF DUPLICATE DEBIT ERR{i:02d}99",
                "debit": Decimal("0.00"), "credit": Decimal("14250.00"), "ref": f"REV{i:02d}99"
            },
            # Customer Collections (AR Invoices)
            {
                "day": 14, "desc": f"INWARD ACH PYMT ORANGE DATA TELECOM INV-2026-{i:02d}01",
                "debit": Decimal("0.00"), "credit": Decimal("92400.00"), "ref": f"DEMO-INV-2026-{i:02d}01",
                "entity": "Orange Data Telecom", "rec_type": "AR_INVOICE"
            },
            {
                "day": 16, "desc": f"INSTAPAY IPN RECEIPT VODAFONE EGYPT INV-2026-{i:02d}02",
                "debit": Decimal("0.00"), "credit": Decimal("48600.00"), "ref": f"DEMO-INV-2026-{i:02d}02",
                "entity": "Vodafone Egypt", "rec_type": "AR_INVOICE"
            },
            {
                "day": 18, "desc": f"FAWRY PAY MERCHANDISE SETTLEMENT INV-2026-{i:02d}03",
                "debit": Decimal("0.00"), "credit": Decimal("35120.00"), "ref": f"DEMO-INV-2026-{i:02d}03",
                "entity": "Fawry Pay", "rec_type": "AR_INVOICE"
            },
            # Supplier Disbursements (AP Bills)
            {
                "day": 20, "desc": f"OUTWARD WIRE SCHNEIDER ELECTRIC BILL-2026-{i:02d}10",
                "debit": Decimal("145000.00"), "credit": Decimal("0.00"), "ref": f"DEMO-BILL-2026-{i:02d}10",
                "entity": "Schneider Electric", "rec_type": "AP_BILL"
            },
            {
                "day": 22, "desc": f"VENDOR DISBURSEMENT BAKER HUGHES EGYPT BILL-2026-{i:02d}20",
                "debit": Decimal("82500.00"), "credit": Decimal("0.00"), "ref": f"DEMO-BILL-2026-{i:02d}20",
                "entity": "Baker Hughes Egypt", "rec_type": "AP_BILL"
            },
            # Corporate Payroll
            {
                "day": 25, "desc": "MONTHLY NET SALARY BULK DISBURSEMENT BATCH 01",
                "debit": Decimal("380000.00"), "credit": Decimal("0.00"), "ref": f"PAYROLL-{i:02d}"
            },
            # Trade Finance LG
            {
                "day": 26, "desc": f"LETTER OF GUARANTEE COMMISSION ISSUANCE LG-2026-{i:02d}50",
                "debit": Decimal("4500.00"), "credit": Decimal("0.00"), "ref": f"DEMO-LG-2026-{i:02d}50",
                "entity": "National Contracting Corp", "rec_type": "LG_COMMISSION"
            },
            # Date Disambiguation Test Line (explicit month token)
            {
                "day": min(27, end_day), "desc": f"BRANCH CASH COLLECTION REF 05FEB2026 POSTING",
                "debit": Decimal("0.00"), "credit": Decimal("25000.00"), "ref": f"DEP-{i:02d}88"
            },
            # Cryptic boundary memo
            {
                "day": min(28, end_day), "desc": f"MISC TRF BNK MEMO 90124809214{i:02d}",
                "debit": Decimal("0.00"), "credit": Decimal("1500.00"), "ref": f"MEMO-{i:02d}99"
            }
        ]

        # Calculate exact running balances
        processed_rows = []
        for r in raw_txns:
            txn_date = datetime(stmt_year, stmt_month, r["day"])
            net_change = r["credit"] - r["debit"]
            current_bal += net_change
            
            processed_rows.append({
                "Date": txn_date.strftime("%d/%m/%Y"),
                "Description": r["desc"],
                "Debit": float(r["debit"]) if r["debit"] > 0 else None,
                "Credit": float(r["credit"]) if r["credit"] > 0 else None,
                "Balance": float(current_bal),
                "Reference": r["ref"]
            })

            # Seed matching ERP record if applicable
            if "rec_type" in r:
                erp_amt = r["credit"] if r["rec_type"] == "AR_INVOICE" else -r["debit"]
                erp_rec = InternalLedgerRecord(
                    company_id=customer_id,
                    record_type=r["rec_type"],
                    reference_number=r["ref"],
                    entity_name=r["entity"],
                    record_date=txn_date,
                    amount=erp_amt,
                    currency="EGP",
                    gl_account="4010010" if r["rec_type"] == "AR_INVOICE" else "2010010",
                    status="OPEN"
                )
                db.add(erp_rec)
                total_erp_records_created += 1

        closing_bal = current_bal

        # ── Build Styled Excel Workbook ──
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Statement"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_header_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        font_sub = Font(name="Calibri", size=10, color="475569")
        font_tbl_head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        font_mono = Font(name="Consolas", size=9.5)

        fill_tbl_head = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_meta = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_stripe = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # 1. Header Metadata block
        ws.merge_cells("A1:F1")
        ws["A1"] = f"{bank_name} — Account Statement"
        ws["A1"].font = font_header_title
        ws["A1"].alignment = Alignment(vertical="center")

        ws["A2"] = "Account Number:"
        ws["B2"] = account_no
        ws["D2"] = "Currency:"
        ws["E2"] = "EGP"

        ws["A3"] = "Statement Period:"
        ws["B3"] = f"{stmt_start.strftime('%d/%m/%Y')} to {stmt_end.strftime('%d/%m/%Y')}"
        ws["D3"] = "Opening Balance:"
        ws["E3"] = float(opening_bal)
        ws["E3"].number_format = "#,##0.00"

        ws["A4"] = "Customer Name:"
        ws["B4"] = "Grow Corporate Client S.A.E."
        ws["D4"] = "Closing Balance:"
        ws["E4"] = float(closing_bal)
        ws["E4"].number_format = "#,##0.00"

        for r_idx in range(2, 5):
            for c_idx in range(1, 7):
                ws.cell(row=r_idx, column=c_idx).font = font_sub

        # 2. Table Column Headers on Row 6
        headers = ["Date", "Description", "Debit", "Credit", "Balance", "Reference"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col_idx, value=h)
            cell.font = font_tbl_head
            cell.fill = fill_tbl_head
            cell.alignment = Alignment(horizontal="right" if h in ["Debit", "Credit", "Balance"] else "left")
            cell.border = border_thin

        # 3. Data Rows starting on Row 7
        for r_idx, row_data in enumerate(processed_rows, 7):
            c_date = ws.cell(row=r_idx, column=1, value=row_data["Date"])
            c_desc = ws.cell(row=r_idx, column=2, value=row_data["Description"])
            c_deb = ws.cell(row=r_idx, column=3, value=row_data["Debit"])
            c_cred = ws.cell(row=r_idx, column=4, value=row_data["Credit"])
            c_bal = ws.cell(row=r_idx, column=5, value=row_data["Balance"])
            c_ref = ws.cell(row=r_idx, column=6, value=row_data["Reference"])

            c_date.font = font_mono
            c_desc.font = font_data
            c_deb.font = font_mono
            c_cred.font = font_mono
            c_bal.font = font_mono
            c_ref.font = font_mono

            c_deb.number_format = "#,##0.00"
            c_cred.number_format = "#,##0.00"
            c_bal.number_format = "#,##0.00"

            c_deb.alignment = Alignment(horizontal="right")
            c_cred.alignment = Alignment(horizontal="right")
            c_bal.alignment = Alignment(horizontal="right")

            is_alt = (r_idx % 2 == 0)
            for c in [c_date, c_desc, c_deb, c_cred, c_bal, c_ref]:
                c.border = border_thin
                if is_alt:
                    c.fill = fill_stripe

        # Column widths
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 52
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 24

        file_name = f"Customer_Demo_{i:02d}_{bank_code}.xlsx"
        file_path = os.path.join(OUTPUT_DIR, file_name)
        wb.save(file_path)
        generated_files.append(file_path)
        print(f"[{i:02d}/20] Generated: {file_name} ({len(processed_rows)} txns, Closing Bal: {closing_bal:,.2f})")

    db.commit()
    db.close()

    print("\n" + "="*70)
    print(f"SUCCESS: Generated 20 Customer Demo Statements in: {OUTPUT_DIR}")
    print(f"Total simulated ERP ledger records created in DB: {total_erp_records_created}")
    print("="*70)

if __name__ == "__main__":
    generate_statements_and_erp_records()
